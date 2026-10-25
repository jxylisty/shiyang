# -*- coding: utf-8 -*-
"""生成方子整理模板 Excel：docs/方子整理模板.xlsx"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")
OUT = Path(__file__).resolve().parent.parent / "docs" / "方子整理模板.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="2E7D5B")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
EXAMPLE_FILL = PatternFill("solid", fgColor="F1F8F4")

wb = Workbook()

# ---------- Sheet 1: 填写说明 ----------
ws = wb.active
ws.title = "填写说明"
notes = [
    ["「食养有据」方子数据整理模板 · 填写说明"],
    [""],
    ["1. 本模板两个业务表：【方子表】一行一个方子；【配料表】一行一种食材（一个方子对应多行配料）。"],
    ["2. 【配料表.食材名】请优先写：药食同源目录正式名（如'枸杞子'不写'枸杞'）或现有药材库名（如'陈皮'不写'橘皮'）；"],
    ["   普通食物（银耳/雪梨/普洱茶等）随便写，后续进 food 表。"],
    ["3. 【类型】只填：药食同源 / 普通食物 / 茶叶 三种，决定成分数据还是营养素数据。"],
    ["4. 【适宜体质】从九种里多选逗号分隔：平和/气虚/阳虚/阴虚/痰湿/湿热/血瘀/气郁/特禀；不确定先留空，由体质模块核对。"],
    ["5. 【场景标签】从以下选并可持续扩充：护嗓润喉/熬夜护眼/换季调理/助眠/消化/控糖轻养/控油祛痘/久坐通气/情绪舒压/经期温养。"],
    ["6. 【禁忌提示】必填。孕妇/儿童/慢病等慎用情况写清楚；全文避免'治疗/疗效'字样，用'调理参考'。"],
    ["7. 【出处】【参考文献】必填：教材写《中医食疗学》(版次/页码)，文献写作者+期刊+年份，网方写来源链接。"],
    ["8. 每条数据来源必须可追溯——这是本产品与普通养生App的核心区别。"],
    ["9. 示例行（绿色底）仅演示格式，正式录入时可删除。"],
    [""],
    ["配方设计小提示（基于 2026-10 重合度盘点）："],
    ["  - 现有成分库已覆盖 74/106 种药食同源物质（302个化合物，187个高入血概率），优先用已覆盖物质组方："],
    ["    常用且数据好：枸杞子、菊花、陈皮、山楂、百合、茯苓、酸枣仁、麦冬、黄芪、大枣、生姜、金银花、蒲公英、薄荷、桑叶、乌梅、罗汉果(缺)、葛根、薏苡仁、桔梗、甘草"],
    ["  - 未覆盖但常用的：莲子、荷叶、决明子、砂仁、龙眼肉、黑芝麻、赤小豆 → 走'普通食物+营养素'或等 TCMSP 补数据。"],
]
for r, row in enumerate(notes, 1):
    ws.cell(r, 1, row[0])
ws.cell(1, 1).font = Font(bold=True, size=14)
ws.column_dimensions["A"].width = 110

# ---------- Sheet 2: 方子表 ----------
ws2 = wb.create_sheet("方子表")
recipe_headers = ["方名", "类型(茶饮/药膳)", "一句话功效", "详细功效描述", "做法", "冲泡或炖煮时长",
                  "适宜体质", "不宜体质", "场景标签", "适用人群描述", "禁忌提示", "出处", "参考文献", "状态"]
example_recipes = [
    ["枸杞菊花明目茶", "茶饮", "清润明目，久看屏幕后的温和护眼参考",
     "枸杞子滋补肝肾、益精明目；菊花清散风热、平肝明目。二者配伍常用于用眼疲劳干涩的日常调理参考。",
     "枸杞子6g、菊花4g，沸水冲泡，加盖焖5分钟，可续水2-3次", "冲泡5分钟",
     "平和,阴虚", "阳虚(畏寒明显者菊花减半)", "熬夜护眼", "长期用眼的学生与办公人群",
     "脾胃虚寒便溏者慎用；不与寒凉药物同服", "《中医食疗学》教材同类方+常用配伍", "待核对",
     "示例"],
    ["陈皮山楂消食饮", "茶饮", "餐后油腻感的传统消食化积参考",
     "陈皮理气健脾、燥湿化痰；山楂消食健胃、行气散瘀。适合聚餐油腻后的温和调理参考。",
     "陈皮3g、炒山楂6g，沸水冲泡焖10分钟，或小火煮10分钟", "冲泡10分钟",
     "痰湿,气郁", "气虚(胃酸过多者慎)", "消化", "饮食油腻、饭后腹胀人群",
     "胃酸过多/空腹不宜；孕妇慎用山楂", "《中医食疗学》教材同类方+常用配伍", "待核对", "示例"],
    ["黄芪大枣温养茶", "茶饮", "换季易感期的传统益气固表调理参考",
     "黄芪补气升阳、益卫固表；大枣补中益气、养血安神。常用于气虚乏力、换季易感人群的日常温养参考。",
     "黄芪5g、大枣2枚(掰开)，沸水冲泡或煮15分钟", "煮15分钟",
     "气虚,阳虚", "湿热,阴虚(上火期停用)", "换季调理", "乏力易感、手脚偏凉人群",
     "感冒发热期停用；湿热壅盛者不宜", "《中医食疗学》教材同类方+常用配伍", "待核对", "示例"],
]
for c, h in enumerate(recipe_headers, 1):
    cell = ws2.cell(1, c, h)
    cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")
for r, row in enumerate(example_recipes, 2):
    for c, v in enumerate(row, 1):
        cell = ws2.cell(r, c, v)
        cell.fill = EXAMPLE_FILL
for c, w in zip(range(1, len(recipe_headers) + 1),
                [16, 12, 26, 40, 34, 12, 14, 16, 12, 20, 24, 26, 26, 8]):
    ws2.column_dimensions[get_column_letter(c)].width = w
ws2.freeze_panes = "A2"

# ---------- Sheet 3: 配料表 ----------
ws3 = wb.create_sheet("配料表")
material_headers = ["方名", "食材名", "类型(药食同源/普通食物/茶叶)", "克数(g)", "是否可选", "备注"]
example_materials = [
    ["枸杞菊花明目茶", "枸杞子", "药食同源", 6, "必选", "明目主料"],
    ["枸杞菊花明目茶", "菊花", "药食同源", 4, "必选", "畏寒者换白菊花减量"],
    ["陈皮山楂消食饮", "陈皮", "药食同源", 3, "必选", ""],
    ["陈皮山楂消食饮", "山楂", "药食同源", 6, "必选", "建议炒制"],
    ["黄芪大枣温养茶", "黄芪", "药食同源", 5, "必选", ""],
    ["黄芪大枣温养茶", "大枣", "药食同源", 6, "必选", "2枚约6g，掰开入水"],
]
for c, h in enumerate(material_headers, 1):
    cell = ws3.cell(1, c, h)
    cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")
for r, row in enumerate(example_materials, 2):
    for c, v in enumerate(row, 1):
        cell = ws3.cell(r, c, v)
        cell.fill = EXAMPLE_FILL
for c, w in zip(range(1, len(material_headers) + 1), [20, 14, 24, 10, 10, 26]):
    ws3.column_dimensions[get_column_letter(c)].width = w
ws3.freeze_panes = "A2"

wb.save(OUT)
print(f"模板已生成: {OUT}")
print("Sheet: 填写说明 / 方子表(3示例) / 配料表(6示例)")
